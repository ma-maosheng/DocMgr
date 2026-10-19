# -*- coding: utf-8 -*-
"""根据 MainWindow 左侧导航生成菜单结构图 PPT。"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "测绘资料管理系统-菜单结构图.pptx"

TEAL = RGBColor(0x39, 0x7E, 0x7A)
TEAL_DARK = RGBColor(0x2A, 0x5F, 0x5C)
NAVY = RGBColor(0x1F, 0x2A, 0x37)
GRAY = RGBColor(0x4A, 0x55, 0x63)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xF4, 0xF7, 0xF7)
CARD = RGBColor(0xFF, 0xFF, 0xFF)
LEAF_BG = RGBColor(0xE8, 0xF3, 0xF2)
GROUP_BG = RGBColor(0xD6, 0xE8, 0xE6)
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

# 与 Views/MainWindow.xaml 左侧 NavMenuPanel 一致
MENU: list[dict] = [
    {
        "name": "历史存档",
        "items": [
            "地形图图件",
            "航摄胶片、像片",
            "其他资料",
            {
                "name": "资料台账与资料处置",
                "items": ["台账", "处置"],
            },
        ],
    },
    {
        "name": "年度资料",
        "items": [
            "项目信息",
            {
                "name": "建档与立档",
                "items": [
                    "建档申请",
                    "建档审批",
                    "资料立档",
                    "存量硬盘直办",
                    "存档文本直办",
                    "立档台账",
                ],
            },
            {
                "name": "迁档",
                "items": [
                    "模拟介质资料迁档",
                    "电子介质资料迁档",
                    "迁档台账",
                ],
            },
            {
                "name": "检索",
                "items": [
                    "资料检索(综合模式)",
                    "电子介质资料检索",
                    "模拟介质资料检索",
                    "检索池",
                ],
            },
            {
                "name": "实物流转",
                "items": [
                    "借出申请",
                    "借出审批",
                    "归还申请",
                    "归还审批",
                    "流转台账",
                ],
            },
            {
                "name": "出入网",
                "items": [
                    "入网申请",
                    "入网审批",
                    "出网申请",
                    "出网审批",
                    "在网数据台账",
                    "在网数据处置",
                    "跨域流转台账",
                ],
            },
            {
                "name": "盘库与离库",
                "items": [
                    "模拟资料盘库",
                    "电子资料盘库",
                    "模拟资料离库处置",
                    "电子资料离库处置",
                ],
            },
        ],
    },
    {
        "name": "库房与介质",
        "items": [
            "库管资料总览",
            {
                "name": "档案柜",
                "items": ["档案柜登记", "开柜查看"],
            },
            {
                "name": "硬盘",
                "items": [
                    "概览",
                    "初始登记",
                    "出库申请",
                    "出库办理",
                    "归还申请",
                    "归还办理",
                    "盘库登记",
                    "离库处置",
                    "硬盘台账",
                ],
            },
            {
                "name": "光盘",
                "items": ["概览", "光盘台账"],
            },
        ],
    },
    {
        "name": "系统设置",
        "items": [
            "个人设置",
            "用户管理",
            "部门设置",
            "角色概览",
            "审核审批",
            "服务器路径设置",
            "逾期设置",
            "数据库操作日志",
            "在线用户与登录记录",
            "待办事项总览",
            "高级数据管理",
        ],
    },
    {
        "name": "帮助",
        "items": ["操作手册"],
    },
]


def set_run(run, text, size=16, bold=False, color=NAVY, font="Microsoft YaHei"):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_bg(slide, color=LIGHT):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    sp_tree = slide.shapes._spTree
    sp = shape._element
    sp_tree.remove(sp)
    sp_tree.insert(2, sp)


def add_top_bar(slide, title: str):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, Inches(0.78))
    bar.fill.solid()
    bar.fill.fore_color.rgb = TEAL
    bar.line.fill.background()
    box = slide.shapes.add_textbox(Inches(0.4), Inches(0.16), Inches(12.4), Inches(0.5))
    tf = box.text_frame
    tf.word_wrap = True
    set_run(tf.paragraphs[0].add_run(), title, size=22, bold=True, color=WHITE)


def add_footer(slide, page: int, total: int):
    box = slide.shapes.add_textbox(Inches(0.4), Inches(7.12), Inches(12.5), Inches(0.28))
    set_run(
        box.text_frame.paragraphs[0].add_run(),
        f"河北省第三测绘院 · 测绘资料管理系统 · 左侧导航菜单结构  |  {page}/{total}",
        size=10,
        color=GRAY,
    )


def add_rounded_card(slide, left, top, width, height, fill, text, size=12, bold=False, color=NAVY):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = TEAL if fill != TEAL else TEAL_DARK
    shape.line.width = Pt(1)
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.08)
    tf.margin_top = Inches(0.04)
    tf.margin_bottom = Inches(0.04)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    set_run(run, text, size=size, bold=bold, color=color)
    return shape


def count_leaves(node) -> int:
    if isinstance(node, str):
        return 1
    return sum(count_leaves(x) for x in node.get("items", []))


def flatten_paths(node, prefix: list[str] | None = None) -> list[str]:
    prefix = prefix or []
    if isinstance(node, str):
        return [" · ".join(prefix + [node])]
    name = node["name"]
    paths: list[str] = []
    for child in node["items"]:
        paths.extend(flatten_paths(child, prefix + [name]))
    return paths


def cover_slide(prs, total: int):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, TEAL_DARK)
    title = slide.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(11.7), Inches(1))
    p = title.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    set_run(p.add_run(), "测绘资料管理系统", size=36, bold=True, color=WHITE)
    sub = slide.shapes.add_textbox(Inches(0.8), Inches(3.3), Inches(11.7), Inches(0.6))
    p2 = sub.text_frame.paragraphs[0]
    p2.alignment = PP_ALIGN.CENTER
    set_run(p2.add_run(), "左侧导航菜单结构图", size=26, bold=True, color=RGBColor(0xD6, 0xFF, 0xF8))
    note = slide.shapes.add_textbox(Inches(0.8), Inches(4.3), Inches(11.7), Inches(0.8))
    p3 = note.text_frame.paragraphs[0]
    p3.alignment = PP_ALIGN.CENTER
    leaf_total = sum(count_leaves(m) for m in MENU)
    set_run(
        p3.add_run(),
        f"河北省第三测绘院  ·  共 {len(MENU)} 个一级模块 / {leaf_total} 个功能入口",
        size=16,
        color=WHITE,
    )
    add_footer(slide, 1, total)


def overview_slide(prs, page: int, total: int):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_top_bar(slide, "一、导航总览（与主窗口左侧菜单一致）")

    cards = [
        ("历史存档", "图件台账与历史离库"),
        ("年度资料", "建档立档→流转→出入网→盘库离库"),
        ("库房与介质", "柜体 · 硬盘 · 光盘"),
        ("系统设置", "用户权限 · 审核 · 运维"),
        ("帮助", "操作手册"),
    ]
    left0 = Inches(0.45)
    top = Inches(1.2)
    w = Inches(2.35)
    h = Inches(1.55)
    gap = Inches(0.18)
    for i, (name, desc) in enumerate(cards):
        x = left0 + (w + gap) * i
        add_rounded_card(slide, x, top, w, Inches(0.55), TEAL, name, size=16, bold=True, color=WHITE)
        body = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, top + Inches(0.62), w, Inches(0.85))
        body.fill.solid()
        body.fill.fore_color.rgb = CARD
        body.line.color.rgb = TEAL
        tf = body.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        set_run(p.add_run(), desc, size=12, color=NAVY)

    # 结构说明
    lines = [
        "层级约定：一级模块（手风琴）→ 二级分组/直达菜单 → 三级功能页",
        "年度资料按生命周期编排：项目信息 → 建档与立档 → 迁档 → 检索 → 实物流转 → 出入网 → 盘库与离库",
        "库房与介质：库管总览 + 档案柜 + 硬盘全流程 + 光盘台账",
        "导航全员可进浏览；写入/办理仍按角色控制（详见「角色概览」）",
    ]
    box = slide.shapes.add_textbox(Inches(0.5), Inches(3.2), Inches(12.3), Inches(3.4))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(10)
        set_run(p.add_run(), f"• {line}", size=15, color=NAVY)
    add_footer(slide, page, total)


def module_tree_slide(prs, page: int, total: int, module: dict):
    """一页展示一个一级模块的树状结构。"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_top_bar(slide, f"二、{module['name']}")

    # 根节点
    add_rounded_card(
        slide,
        Inches(0.4),
        Inches(1.05),
        Inches(2.4),
        Inches(0.48),
        TEAL,
        module["name"],
        size=14,
        bold=True,
        color=WHITE,
    )

    items = module["items"]
    # 二级列布局
    col_count = min(len(items), 4 if module["name"] != "年度资料" else 4)
    if module["name"] == "年度资料":
        # 两行：项目+前3组 / 后3组
        _layout_yearly(slide, items)
    elif module["name"] == "系统设置":
        _layout_flat_chips(slide, items, start_top=1.75)
    else:
        _layout_columns(slide, items, start_top=1.75, max_cols=col_count)

    add_footer(slide, page, total)


def _layout_columns(slide, items, start_top=1.75, max_cols=4):
    usable_w = 12.4
    cols = min(len(items), max_cols)
    col_w = usable_w / cols
    chip_h = 0.36
    for i, item in enumerate(items):
        col = i % cols
        row_base = start_top + (i // cols) * 2.8
        x = 0.45 + col * col_w
        if isinstance(item, str):
            add_rounded_card(
                slide,
                Inches(x),
                Inches(row_base),
                Inches(col_w - 0.15),
                Inches(chip_h),
                LEAF_BG,
                item,
                size=11,
                bold=True,
            )
        else:
            add_rounded_card(
                slide,
                Inches(x),
                Inches(row_base),
                Inches(col_w - 0.15),
                Inches(chip_h),
                GROUP_BG,
                item["name"],
                size=11,
                bold=True,
            )
            for j, leaf in enumerate(item["items"]):
                add_rounded_card(
                    slide,
                    Inches(x + 0.08),
                    Inches(row_base + 0.42 + j * (chip_h + 0.06)),
                    Inches(col_w - 0.28),
                    Inches(chip_h),
                    CARD,
                    leaf,
                    size=10,
                )


def _layout_yearly(slide, items):
    """年度资料：7 项分两行排布。"""
    # 第一行：项目信息 + 建档与立档 + 迁档 + 检索
    row1 = items[:4]
    row2 = items[4:]
    _draw_yearly_row(slide, row1, top=1.7)
    _draw_yearly_row(slide, row2, top=4.35)


def _draw_yearly_row(slide, row_items, top: float):
    cols = len(row_items)
    col_w = 12.4 / cols
    chip_h = 0.32
    for i, item in enumerate(row_items):
        x = 0.45 + i * col_w
        if isinstance(item, str):
            add_rounded_card(
                slide,
                Inches(x),
                Inches(top),
                Inches(col_w - 0.12),
                Inches(chip_h + 0.05),
                GROUP_BG,
                item,
                size=11,
                bold=True,
            )
            continue
        add_rounded_card(
            slide,
            Inches(x),
            Inches(top),
            Inches(col_w - 0.12),
            Inches(chip_h),
            GROUP_BG,
            item["name"],
            size=11,
            bold=True,
        )
        for j, leaf in enumerate(item["items"]):
            add_rounded_card(
                slide,
                Inches(x + 0.06),
                Inches(top + 0.38 + j * (chip_h + 0.04)),
                Inches(col_w - 0.22),
                Inches(chip_h),
                CARD,
                leaf,
                size=9,
            )


def _layout_flat_chips(slide, items, start_top=1.75):
    """系统设置：扁平芯片网格。"""
    cols = 3
    chip_w = 3.9
    chip_h = 0.42
    gap_x = 0.2
    gap_y = 0.14
    for i, item in enumerate(items):
        name = item if isinstance(item, str) else item["name"]
        col = i % cols
        row = i // cols
        x = 0.55 + col * (chip_w + gap_x)
        y = start_top + row * (chip_h + gap_y)
        add_rounded_card(slide, Inches(x), Inches(y), Inches(chip_w), Inches(chip_h), LEAF_BG, name, size=13, bold=True)


def inventory_slide(prs, page: int, total: int):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_top_bar(slide, "三、功能入口清单（完整路径）")

    all_paths: list[str] = []
    for m in MENU:
        all_paths.extend(flatten_paths(m))

    mid = (len(all_paths) + 1) // 2
    left_paths = all_paths[:mid]
    right_paths = all_paths[mid:]

    def fill_column(left, paths):
        box = slide.shapes.add_textbox(Inches(left), Inches(1.0), Inches(6.1), Inches(5.9))
        tf = box.text_frame
        tf.word_wrap = True
        for i, path in enumerate(paths):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(2)
            set_run(p.add_run(), path, size=10, color=NAVY)

    fill_column(0.4, left_paths)
    fill_column(6.8, right_paths)
    add_footer(slide, page, total)


def build():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    # 封面 + 总览 + 5 模块 + 清单
    total = 1 + 1 + len(MENU) + 1
    cover_slide(prs, total)
    overview_slide(prs, 2, total)
    for i, module in enumerate(MENU):
        module_tree_slide(prs, 3 + i, total, module)
    inventory_slide(prs, total, total)

    prs.save(OUT)
    print(f"已生成: {OUT}")
    print(f"共 {total} 页")


if __name__ == "__main__":
    build()
